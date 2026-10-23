"""Genera el oficio de remisión en Word a partir de la plantilla institucional.

La plantilla aporta el membrete (logos, encabezado y pie de página).  Aquí se reemplazan
la fecha, la destinataria, el asunto, el cuerpo con el valor total, el bloque de firma y
se inserta la tabla del informe como imagen, en página nueva.
"""
import copy
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm
from docx.text.paragraph import Paragraph

from . import imagen_tabla


class _Contexto(dict):
    """Deja intactas las llaves que no se hayan definido en el contexto."""

    def __missing__(self, clave):
        return "{" + clave + "}"


def _poner_texto(parrafo, texto, negrita=None):
    """Escribe el texto conservando el formato; si `negrita` no es None, lo fija."""
    if parrafo.runs:
        parrafo.runs[0].text = texto
        for run in parrafo.runs[1:]:
            run.text = ""
        if negrita is not None:
            for run in parrafo.runs:
                run.bold = negrita
    else:
        run = parrafo.add_run(texto)
        if negrita is not None:
            run.bold = negrita


def _buscar_parrafo(doc, patron, desde=0):
    for indice in range(desde, len(doc.paragraphs)):
        if re.search(patron, doc.paragraphs[indice].text, re.IGNORECASE):
            return indice
    return None


def _siguiente_con_texto(doc, desde):
    for indice in range(desde, len(doc.paragraphs)):
        if doc.paragraphs[indice].text.strip():
            return indice
    return None


def _escribir_bloque(doc, indices, textos):
    """Escribe un bloque de líneas; cada línea puede ser texto o (texto, negrita)."""
    for posicion, indice in enumerate(indices):
        valor = textos[posicion] if posicion < len(textos) else ""
        negrita = None
        if isinstance(valor, tuple):
            valor, negrita = valor
        _poner_texto(doc.paragraphs[indice], valor or "", negrita)
    if indices and len(textos) > len(indices):
        ultimo = doc.paragraphs[indices[-1]]
        for valor in textos[len(indices):]:
            negrita = None
            if isinstance(valor, tuple):
                valor, negrita = valor
            nuevo = copy.deepcopy(ultimo._p)
            ultimo._p.addnext(nuevo)
            ultimo = Paragraph(nuevo, ultimo._parent)
            _poner_texto(ultimo, valor or "", negrita)


def _recortar_vacios_finales(doc):
    while doc.paragraphs and not doc.paragraphs[-1].text.strip():
        elemento = doc.paragraphs[-1]._p
        elemento.getparent().remove(elemento)


def _lineas_destinatarios(oficio):
    """Líneas del bloque de destinatarios, con un renglón en blanco entre cada funcionario.

    La segunda línea de cada bloque es el nombre del funcionario y va en negrilla.
    Acepta `destinatarios` (lista de bloques) y, por compatibilidad, `destinatario`.
    """
    bloques = oficio.get("destinatarios")
    if not bloques:
        bloques = [oficio.get("destinatario") or []]
    if isinstance(bloques, list) and bloques and isinstance(bloques[0], str):
        bloques = [bloques]
    lineas = []
    for bloque in bloques:
        if lineas:
            lineas.append(("", False))   # renglón en blanco entre funcionarios
        for posicion, linea in enumerate(bloque):
            lineas.append((linea, posicion == 1))   # el nombre va en negrilla
    return lineas


def medidas_utiles(doc):
    """Ancho y alto útiles (cm) de la última sección del documento."""
    seccion = doc.sections[-1]
    ancho = seccion.page_width.cm - seccion.left_margin.cm - seccion.right_margin.cm
    alto = seccion.page_height.cm - seccion.top_margin.cm - seccion.bottom_margin.cm
    return ancho, alto


def generar(ruta_plantilla, ruta_salida, cfg, contexto, informe, carpeta_imagenes=None):
    """Devuelve {'ruta', 'paginas_tabla', 'imagenes'}."""
    if not os.path.exists(ruta_plantilla):
        raise FileNotFoundError(f"No se encontró la plantilla del oficio: {ruta_plantilla}")

    doc = Document(ruta_plantilla)
    oficio = cfg["oficio"]
    ctx = _Contexto(contexto)

    # 1. Lugar y fecha
    indice = _buscar_parrafo(doc, r"^\s*" + re.escape(str(oficio.get("ciudad_fecha") or "Soledad")) + r"\s*,")
    if indice is None:
        indice = _buscar_parrafo(doc, r"^\s*[A-ZÁÉÍÓÚ][\wáéíóúñ.]*\s*,\s*\w+\s+\d{1,2}\s+de\s+\d{4}\s*$")
    if indice is not None:
        _poner_texto(doc.paragraphs[indice], f"{oficio.get('ciudad_fecha', 'Soledad')}, {contexto['FECHA']}")

    # 2. Destinatarios (uno o varios bloques, separados por una línea en blanco)
    inicio = _buscar_parrafo(doc, r"^\s*(Doctora?|Se[ñn]ora?|Sr\.|Sra\.)\s*$")
    indice_ref = _buscar_parrafo(doc, r"^\s*Ref\s*[:.]")
    if inicio is not None:
        limite = indice_ref if indice_ref is not None and indice_ref > inicio else len(doc.paragraphs)
        indices = [i for i in range(inicio, limite) if doc.paragraphs[i].text.strip()]
        _escribir_bloque(doc, indices, _lineas_destinatarios(oficio))

    # 3. Asunto (se vuelve a buscar: insertar destinatarios desplaza los párrafos)
    indice_ref = _buscar_parrafo(doc, r"^\s*Ref\s*[:.]")
    if indice_ref is not None:
        _poner_texto(doc.paragraphs[indice_ref], "Ref.: " + str(oficio["asunto"]).format_map(ctx))

    # 4. Cuerpo (con el valor total de la tabla)
    indice_cuerpo = _buscar_parrafo(doc, r"por valor de")
    if indice_cuerpo is not None:
        _poner_texto(doc.paragraphs[indice_cuerpo], str(oficio["cuerpo"]).format_map(ctx))

    # 5. Bloque de firma
    saludo = _buscar_parrafo(doc, r"^\s*Cordialmente\s*,?\s*$")
    if saludo is not None:
        indices, siguiente = [], saludo + 1
        while siguiente is not None and len(indices) < len(oficio["remitente"]):
            if re.search(r"^\s*(Proyect[oó]|Revis[oó])\s*:", doc.paragraphs[siguiente].text):
                break
            if doc.paragraphs[siguiente].text.strip():
                indices.append(siguiente)
            siguiente = siguiente + 1 if siguiente + 1 < len(doc.paragraphs) else None
        _escribir_bloque(doc, indices, oficio["remitente"])

    # 6. Proyectó / Revisó
    for clave, patron in (("proyecto", r"^\s*Proyect[oó]\s*:"), ("reviso", r"^\s*Revis[oó]\s*:")):
        indice = _buscar_parrafo(doc, patron)
        if indice is None:
            continue
        indices = [indice]
        siguiente = _buscar_parrafo(doc, r"\S", desde=indice + 1)
        if siguiente is not None and not re.search(r"^\s*(Proyect[oó]|Revis[oó])\s*:", doc.paragraphs[siguiente].text):
            indices.append(siguiente)
        textos = [str(t).replace("{NOMBRE}", str(oficio.get(f"{clave}_nombre") or ""))
                          .replace("{CARGO}", str(oficio.get(f"{clave}_cargo") or ""))
                  for t in oficio[clave]]
        _escribir_bloque(doc, indices, textos)

    # 7. Tabla del informe como imagen
    imagenes, paginas = [], []
    if oficio.get("insertar_tabla", True) and informe and informe.get("filas"):
        ancho_util, alto_util = medidas_utiles(doc)
        ancho_tabla = float(oficio.get("ancho_tabla_cm") or 0) or max(10.0, ancho_util - 0.2)
        ancho_tabla = min(ancho_tabla, ancho_util)
        alto_tabla = max(5.0, alto_util - 0.4)
        paginas = imagen_tabla.construir(informe["titulo"], informe["encabezados"], informe["filas"],
                                         informe["totales"], informe["tipos"], ancho_tabla, alto_tabla, cfg["tabla"])
        _recortar_vacios_finales(doc)
        if oficio.get("pagina_nueva_para_tabla", True):
            doc.add_page_break()
        carpeta_imagenes = carpeta_imagenes or os.path.dirname(ruta_salida)
        os.makedirs(carpeta_imagenes, exist_ok=True)
        for numero, pagina in enumerate(paginas, start=1):
            archivo = os.path.join(
                carpeta_imagenes,
                f"tabla {contexto['MES_TITULO']} {contexto['ANIO']} - pagina {numero}.png",
            )
            pagina.imagen.save(archivo, "PNG", dpi=(int(cfg["tabla"].get("dpi", 300)),) * 2)
            imagenes.append(archivo)
            parrafo = doc.add_paragraph()
            parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            parrafo.add_run().add_picture(archivo, width=Cm(ancho_tabla))
            if numero < len(paginas):
                doc.add_page_break()

    os.makedirs(os.path.dirname(ruta_salida), exist_ok=True)
    doc.save(ruta_salida)
    return {"ruta": ruta_salida, "paginas_tabla": len(paginas), "imagenes": imagenes}


def a_pdf(ruta_docx, ruta_pdf):
    """Convierte el oficio a PDF usando Word. Devuelve la ruta del PDF."""
    import pythoncom
    import win32com.client

    ruta_docx = os.path.abspath(ruta_docx)
    ruta_pdf = os.path.abspath(ruta_pdf)
    os.makedirs(os.path.dirname(ruta_pdf), exist_ok=True)
    pythoncom.CoInitialize()
    aplicacion = documento = None
    try:
        aplicacion = win32com.client.DispatchEx("Word.Application")
        aplicacion.Visible = False
        aplicacion.DisplayAlerts = 0
        documento = aplicacion.Documents.Open(ruta_docx, ReadOnly=True)
        documento.ExportAsFixedFormat(ruta_pdf, 17)
    except Exception as error:
        raise RuntimeError(
            "No se pudo convertir el oficio a PDF. Verifica que Microsoft Word esté instalado "
            f"({error})"
        )
    finally:
        try:
            if documento is not None:
                documento.Close(False)
        except Exception:
            pass
        try:
            if aplicacion is not None:
                aplicacion.Quit()
        except Exception:
            pass
        try:
            pythoncom.CoUninitialize()
        except Exception:
            pass
    return ruta_pdf
