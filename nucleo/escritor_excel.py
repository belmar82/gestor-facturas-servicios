"""Escribe el archivo Excel del informe: hoja TABLA DE INFORME, hoja DETALLE y SIN COINCIDENCIA.

La hoja TABLA DE INFORME (la que se publica) lleva el título combinado y centrado en la
primera fila, tal como aparece en el oficio; los encabezados van en la fila siguiente.
Las hojas DETALLE y SIN COINCIDENCIA conservan los encabezados en la fila 1 para poder
usarlas con BUSCARV sin ajustes.

La hoja TABLA DE INFORME queda con **fórmulas**, no con números fijos:
  * Facturación    = SUMA(DETALLE!Acueducto;DETALLE!Alcantarillado;DETALLE!Aseo;DETALLE!Otros conceptos)
  * Total deuda    = el "Total deuda" de DETALLE, solo si el consumo es ESTIMADO
  * Valor a Pagar  = Facturación + Total deuda - Aseo - Intereses x Mora
  * Totales        = SUMA de cada columna
Así, si alguien corrige un dato en DETALLE, la tabla se actualiza sola.

Todas las hojas quedan listas para imprimir: apaisadas, todo el ancho en una página y con
el título y los encabezados repetidos en cada página.
"""
import os

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

MONEDA = '#,##0'
ENTERO = '#,##0'
CENTRADO = Alignment(horizontal="center", vertical="center")
DERECHA = Alignment(horizontal="right", vertical="center")
IZQUIERDA = Alignment(horizontal="left", vertical="center")

CABECERA = PatternFill("solid", fgColor="1F4E79")
CABECERA_TEXTO = Font(bold=True, color="FFFFFF")
TOTAL_FILL = PatternFill("solid", fgColor="CDD9E5")
BORDE = Border(*[Side(style="thin", color="9DB2C6")] * 4)
TITULO_TEXTO = Font(bold=True, size=12, color="1F4E79")
TITULO_ALTO = 26
# el cero (o el vacío) no se muestra: sirve para la columna 'Total deuda' de las facturas reales
MONEDA_OCULTA = '#,##0;-#,##0;""'


def letra(indice):
    return get_column_letter(indice)


def columna(titulos, fragmento):
    """Número de la columna cuyo encabezado contiene el texto indicado."""
    for indice, titulo in enumerate(titulos, start=1):
        if fragmento.lower() in str(titulo).lower():
            return indice
    return None


def _anchos(plana, titulos, filas):
    for indice, titulo in enumerate(titulos, start=1):
        largo = len(str(titulo))
        for fila in filas:
            valor = fila[indice - 1]
            if valor is None:
                continue
            if isinstance(valor, (int, float)):
                texto = f"{int(valor):,}".replace(",", ".")
            else:
                texto = str(valor)
            largo = max(largo, len(texto))
        plana.column_dimensions[letra(indice)].width = min(max(largo + 3, 10), 46)


def _preparar_impresion(plana, fila_cabecera):
    plana.page_setup.orientation = "landscape"
    plana.page_setup.paperSize = plana.PAPERSIZE_A4
    plana.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    plana.page_setup.fitToWidth = 1
    plana.page_setup.fitToHeight = 0
    plana.print_title_rows = f"1:{fila_cabecera}"
    plana.print_options.horizontalCentered = True
    plana.page_margins.left = plana.page_margins.right = 0.3
    plana.page_margins.top = plana.page_margins.bottom = 0.4


def _escribir_hoja(libro, nombre, titulos, filas, tipos, totales=None, titulo=None,
                   formulas_filas=None, formulas_totales=None, desfase_filas=0):
    if nombre in libro.sheetnames:
        del libro[nombre]
    plana = libro.create_sheet(nombre)

    fila_cabecera = 1
    if titulo:
        # título en una sola celda combinada y centrada, como en el oficio
        plana.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(titulos))
        celda = plana.cell(row=1, column=1, value=titulo)
        celda.font = TITULO_TEXTO
        celda.alignment = CENTRADO
        for columna_ in range(1, len(titulos) + 1):
            plana.cell(row=1, column=columna_).border = BORDE
        plana.row_dimensions[1].height = TITULO_ALTO
        fila_cabecera = 2

    for indice, titulo_col in enumerate(titulos, start=1):
        celda = plana.cell(row=fila_cabecera, column=indice, value=titulo_col)
        celda.fill = CABECERA
        celda.font = CABECERA_TEXTO
        celda.alignment = CENTRADO if indice < 3 else IZQUIERDA
        celda.border = BORDE

    for numero, fila in enumerate(filas, start=fila_cabecera + 1):
        for indice, valor in enumerate(fila, start=1):
            celda = plana.cell(row=numero, column=indice, value=valor)
            celda.border = BORDE
            tipo = tipos[indice - 1]
            if tipo in ("moneda", "entero", "moneda_oculta"):
                celda.number_format = MONEDA_OCULTA if tipo == "moneda_oculta" else MONEDA
                celda.alignment = DERECHA
            elif tipo in ("poliza", "item", "factura"):
                celda.number_format = '0'
                celda.alignment = CENTRADO
            elif tipo == "centrado":
                celda.alignment = CENTRADO
            else:
                celda.alignment = IZQUIERDA

    # fórmulas de cada fila (por ejemplo Facturación, Total deuda y Valor a Pagar)
    for indice, plantilla in (formulas_filas or {}).items():
        for numero in range(fila_cabecera + 1, fila_cabecera + 1 + len(filas)):
            celda = plana.cell(row=numero, column=indice)
            celda.value = plantilla.replace("{f}", str(numero)).replace("{det}", str(numero - desfase_filas))

    if totales:
        numero = fila_cabecera + len(filas) + 1
        for indice, valor in enumerate(totales, start=1):
            celda = plana.cell(row=numero, column=indice, value=valor)
            celda.fill = TOTAL_FILL
            celda.font = Font(bold=True)
            celda.border = BORDE
            tipo = tipos[indice - 1]
            celda.number_format = MONEDA if tipo in ("moneda", "entero", "moneda_oculta") else "General"
            celda.alignment = DERECHA if tipo in ("moneda", "entero", "moneda_oculta") else CENTRADO
        # totales con fórmula
        for indice, plantilla in (formulas_totales or {}).items():
            celda = plana.cell(row=numero, column=indice)
            celda.value = (plantilla.replace("{primera}", str(fila_cabecera + 1))
                           .replace("{ultima}", str(fila_cabecera + len(filas)))
                           .replace("{col}", letra(indice)))

    plana.freeze_panes = plana.cell(row=fila_cabecera + 1, column=1).coordinate
    if filas:
        plana.auto_filter.ref = (f"A{fila_cabecera}:"
                                 f"{letra(len(titulos))}{fila_cabecera + len(filas)}")
    _anchos(plana, titulos, filas)
    _preparar_impresion(plana, fila_cabecera)
    return plana


def _formulas_informe(cab_inf, cab_det, hoja_detalle, cfg):
    """Fórmulas de la hoja TABLA DE INFORME (se buscan las columnas por su encabezado)."""
    col_facturacion = columna(cab_inf, "facturación")
    col_valor = columna(cab_inf, "valor a pagar")
    col_aseo = columna(cab_inf, "aseo")
    col_consumo = columna(cab_inf, "consumo")
    col_intereses = columna(cab_inf, "intereses")
    col_deuda = columna(cab_inf, "total deuda")
    col_metodo = columna(cab_inf, "real/estimado")

    det_acu = columna(cab_det, "servicio de acueducto")
    det_alc = columna(cab_det, "servicio de alcantarillado")
    det_aseo = columna(cab_det, "servicio de aseo")
    det_otros = columna(cab_det, "servicio otros conceptos")
    det_cobros = columna(cab_det, "otros cobros")
    det_deuda = columna(cab_det, "total deuda")

    formulas_filas = {}

    if col_facturacion and det_otros:
        sumandos = [det_acu, det_alc, det_aseo, det_otros]
        if cfg["general"].get("incluir_otros_cobros_en_facturacion", False):
            sumandos.append(det_cobros)
        formulas_filas[col_facturacion] = "=SUM(" + ",".join(
            f"'{hoja_detalle}'!{letra(col)}" + "{det}" for col in sumandos if col) + ")"

    # Total deuda: solo se trae de DETALLE cuando el consumo es ESTIMADO.
    # Devuelve 0 (no texto) para que la suma del Valor a Pagar no dé error; el cero
    # se ve vacío por el formato de número de la columna.
    if col_deuda and col_metodo and det_deuda:
        formulas_filas[col_deuda] = ('=IF(' + letra(col_metodo) + '{f}="ESTIMADO",'
                                     + f"'{hoja_detalle}'!{letra(det_deuda)}" + '{det},0)')

    # Valor a Pagar = Facturación + Total deuda - Aseo - Intereses x Mora
    if col_valor and col_facturacion and col_aseo:
        partes = [letra(col_facturacion) + "{f}"]
        if col_deuda:
            partes.append("+" + letra(col_deuda) + "{f}")
        partes.append("-" + letra(col_aseo) + "{f}")
        if col_intereses:
            partes.append("-" + letra(col_intereses) + "{f}")
        formulas_filas[col_valor] = "=" + "".join(partes)

    formulas_totales = {}
    for col in (col_consumo, col_facturacion, col_aseo, col_intereses, col_deuda, col_valor):
        if col:
            formulas_totales[col] = f"=SUM({letra(col)}" + "{primera}:" + f"{letra(col)}" + "{ultima})"

    return formulas_filas, formulas_totales


def _formulas_detalle(cab_det):
    """'Total factura a pagar' = suma de los seis conceptos menos «otros cobros».

    Los «otros cobros» (por ejemplo el seguro de vida Mapfre) no se cobran, así que se
    restan de la suma.  Si en otro periodo un concepto de esa columna sí se cobra, basta
    con quitar esa resta.
    """
    col_pagar = columna(cab_det, "total factura a pagar")
    sumandos = [columna(cab_det, nombre) for nombre in
                ("servicio de acueducto", "servicio de alcantarillado", "servicio de aseo",
                 "servicio otros conceptos", "otros cobros", "total deuda")]
    col_cobros = columna(cab_det, "otros cobros")
    if col_pagar and all(sumandos) and col_cobros:
        return {col_pagar: (f"=SUM({letra(min(sumandos))}" + "{f}:"
                            + f"{letra(max(sumandos))}" + "{f})"
                            + f"-{letra(col_cobros)}" + "{f}")}
    return {}


def escribir(ruta, detalle, informe, sin_coincidencia, cfg):
    """Crea el libro con las hojas TABLA DE INFORME, DETALLE y SIN COINCIDENCIA."""
    nombres = cfg["informe"]
    hoja_informe = nombres["hoja_informe"]
    hoja_detalle = nombres["hoja_detalle"]

    libro = openpyxl.Workbook()
    libro.remove(libro.active)

    cab_inf, cab_det = informe["encabezados"], detalle["encabezados"]
    formulas_filas, formulas_totales = _formulas_informe(cab_inf, cab_det, hoja_detalle, cfg)
    if informe.get("total_deuda_en_cero"):
        # desde el periodo de corte la columna «Total deuda» queda en 0, sin fórmula
        formulas_filas.pop(columna(cab_inf, "total deuda"), None)

    _escribir_hoja(libro, hoja_informe, cab_inf, informe["filas"], informe["tipos"],
                   informe["totales"], titulo=informe.get("titulo"),
                   formulas_filas=formulas_filas, formulas_totales=formulas_totales,
                   desfase_filas=1)          # TABLA DE INFORME empieza una fila más abajo (título)
    _escribir_hoja(libro, hoja_detalle, cab_det, detalle["filas"], detalle["tipos"],
                   formulas_filas=_formulas_detalle(cab_det))
    if sin_coincidencia and sin_coincidencia["filas"]:
        _escribir_hoja(libro, nombres["hoja_sin_coincidencia"], sin_coincidencia["encabezados"],
                       sin_coincidencia["filas"], sin_coincidencia["tipos"],
                       sin_coincidencia.get("totales"))

    # que Excel recalcule las fórmulas al abrir el archivo
    libro.calculation.fullCalcOnLoad = True

    os.makedirs(os.path.dirname(ruta), exist_ok=True)
    libro.save(ruta)
    return ruta
