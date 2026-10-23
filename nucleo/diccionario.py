"""Lectura del diccionario de pólizas (catálogo de Instituciones Educativas).

El diccionario es la guía de la aplicación: relaciona la póliza con la Institución
Educativa.  Se puede reemplazar por otro archivo o por otro mes sin tocar el programa.
"""
import os
import re

import openpyxl

EQUIVALENCIAS = {
    "poliza": ("póliza", "poliza", "no. poliza", "poliza no"),
    "item": ("item", "ítem", "no", "n°"),
    "ie": ("institución educativa", "institucion educativa", "institucion", "i.e", "ie"),
    "nombre": ("nombre", "cliente", "nombre del cliente"),
    "direccion": ("dirección", "direccion"),
}


def _normalizar(valor):
    texto = str(valor or "").strip().lower().replace("\u00a0", " ")
    texto = re.sub(r"\s+", " ", texto)
    return texto


def _columna(cabeceras, llaves):
    for indice, titulo in enumerate(cabeceras):
        limpio = _normalizar(titulo)
        for llave in llaves:
            if limpio == llave or limpio.startswith(llave):
                return indice
    return None


def elegir_hoja(libro, hoja=None):
    if hoja and hoja in libro.sheetnames:
        return libro[hoja]
    for nombre in libro.sheetnames:
        plana = libro[nombre]
        for fila in plana.iter_rows(min_row=1, max_row=5, values_only=True):
            if any("poliza" in _normalizar(c) or "póliza" in _normalizar(c) for c in fila if c):
                return plana
    return libro[libro.sheetnames[0]]


def leer(ruta, hoja=None):
    """Devuelve {'por_poliza': {...}, 'orden': [pólizas...], 'hoja': nombre}."""
    if not ruta or not os.path.exists(ruta):
        raise FileNotFoundError(f"No se encontró el diccionario de pólizas: {ruta}")
    libro = openpyxl.load_workbook(ruta, data_only=True)
    plana = elegir_hoja(libro, hoja)

    fila_cabecera, cabeceras = None, None
    for numero, fila in enumerate(plana.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
        if any("poliza" in _normalizar(c) or "póliza" in _normalizar(c) for c in fila if c):
            fila_cabecera, cabeceras = numero, list(fila)
            break
    if fila_cabecera is None:
        raise ValueError("El diccionario no tiene una columna 'Póliza'.")

    col = {clave: _columna(cabeceras, llaves) for clave, llaves in EQUIVALENCIAS.items()}
    if col["poliza"] is None:
        raise ValueError("El diccionario no tiene una columna 'Póliza'.")

    por_poliza, orden = {}, []
    for numero, fila in enumerate(plana.iter_rows(min_row=fila_cabecera + 1, values_only=True), start=fila_cabecera + 1):
        if not fila or all(c is None or str(c).strip() == "" for c in fila):
            continue
        crudo = fila[col["poliza"]] if col["poliza"] < len(fila) else None
        if crudo is None:
            continue
        digitos = re.sub(r"[^\d]", "", str(crudo))
        if not digitos:
            continue
        poliza = int(digitos)

        def valor(clave):
            indice = col[clave]
            if indice is None or indice >= len(fila):
                return None
            dato = fila[indice]
            return str(dato).strip() if isinstance(dato, str) else dato

        por_poliza[poliza] = {
            "poliza": poliza,
            "item": valor("item"),
            "ie": valor("ie"),
            "nombre": valor("nombre"),
            "direccion": valor("direccion"),
            "fila": numero,
        }
        orden.append(poliza)

    return {"ruta": ruta, "hoja": plana.title, "por_poliza": por_poliza, "orden": orden}
